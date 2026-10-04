import ast, io, types
from pathlib import Path
from openpyxl import Workbook, load_workbook
import pandas as pd
src=Path('/mnt/data/coe_dummy_manager.py').read_text()
tree=ast.parse(src)
want={'norm','canon','find_header','find_column','source_records','build_mapping','unique_name','copy_cell','sort_sheet_rows','process_targets'}
mod=ast.Module(body=[n for n in tree.body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name in want],type_ignores=[])
ns={'ALIASES_DUMMY':('dummy','dummy no','dummy number','dummy reg no','dummy register no'),'ALIASES_REG':('reg no','reg.no','registration no','registration number','register no','register number','university no','roll no'),'pd':pd,'io':io,'Path':Path,'re':__import__('re'),'load_workbook':load_workbook,'copy':__import__('copy').copy,'Translator':__import__('openpyxl.formula.translate',fromlist=['Translator']).Translator}
exec(compile(mod,'coe_dummy_manager.py','exec'),ns)
# Source workbook: header row 7; two worksheets and conflicting dummy 999
wb=Workbook(); ws=wb.active; ws.title='Day1'
for r in range(1,7): ws.cell(r,1,f'Note {r}')
ws.cell(7,1,'Dummy No.'); ws.cell(7,2,'Reg.No')
ws.append(['D001','R001']); ws.append(['D002','R002']); ws.append(['D999','R003'])
ws2=wb.create_sheet('Day2'); ws2.append(['Dummy Number','Registration Number']); ws2.append(['D999','R004']); ws2.append(['D003','R003'])
buf=io.BytesIO(); wb.save(buf); buf.seek(0); buf.name='same.xlsx'
records,issues=ns['source_records']([buf]); mapping,conflicts,evidence=ns['build_mapping'](records)
assert mapping=={'D001':'R001','D002':'R002','D003':'R003'} and conflicts=={'D999':['R003','R004']},(mapping,conflicts,issues)
# Target workbook header on row 7; formulas; matched, conflict, unmatched; test sort
w=Workbook(); t=w.active; t.title='Marks'
for r in range(1,7): t.cell(r,1,f'heading {r}')
t.cell(7,1,'Dummy No.'); t.cell(7,2,'Score'); t.cell(7,3,'Calc')
t.append(['D002',20,'=B8*2']); t.append(['D999',30,'=B9*2']); t.append(['D404',40,'=B10*2']); t.append(['D001',10,'=B11*2'])
target=io.BytesIO(); w.save(target); target.seek(0); target.name='same.xlsx'
outputs,audit,totals,errors=ns['process_targets']([target,target],mapping,conflicts,'Reg.No','Reg.No Edited',True,False)
assert not errors,errors
assert len(outputs)==2 and outputs[0][0]!=outputs[1][0], [x[0] for x in outputs]
assert totals['rows']==8 and totals['matched']==4 and totals['conflicts']==2 and totals['unmatched']==2,totals
for name,data in outputs:
    out=load_workbook(io.BytesIO(data),data_only=False); sh=out['Marks']
    assert sh.cell(7,1).value=='Dummy No.'
    # sorted registrations: blanks (conflict/unmatched) first per sort function? blank-last key puts blank last
    vals=[sh.cell(r,5).value for r in range(8,12)]
    assert vals==['R001','R002',None,None],vals
    # Formula references translated with row movement: D001 originally row 11 now row 8 -> =B8*2
    assert sh['C8'].value=='=B8*2',sh['C8'].value
    assert sh['C9'].value=='=B9*2',sh['C9'].value
print('PASS: multi-source aggregation, non-row-1 headers, conflict/unmatched detection, duplicate filename protection, formula-aware sort, output reopen')
print('TOTALS:',totals)
